"""PDF/DOCX parsing utilities."""

from pathlib import Path

import pypdf
from docx import Document


def extract_text_from_pdf(file_path: Path | str) -> str:
    """Extract text content from a PDF file."""
    file_path = Path(file_path)
    text_parts = []

    with open(file_path, "rb") as f:
        reader = pypdf.PdfReader(f)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)

    return "\n\n".join(text_parts)


def extract_text_from_pdf_bytes(content: bytes) -> str:
    """Extract text content from PDF bytes."""
    import io

    text_parts = []
    reader = pypdf.PdfReader(io.BytesIO(content))

    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)

    return "\n\n".join(text_parts)


def extract_text_from_docx(file_path: Path | str) -> str:
    """Extract text content from a DOCX file."""
    file_path = Path(file_path)
    doc = Document(file_path)
    text_parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)

    return "\n\n".join(text_parts)


def extract_text_from_docx_bytes(content: bytes) -> str:
    """Extract text content from DOCX bytes."""
    import io

    doc = Document(io.BytesIO(content))
    text_parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)

    return "\n\n".join(text_parts)


def extract_text(file_path: Path | str) -> str:
    """Extract text from a document based on its extension."""
    file_path = Path(file_path)
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)
    elif extension in (".docx", ".doc"):
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {extension}")


def extract_text_from_bytes(content: bytes, file_type: str) -> str:
    """Extract text from document bytes based on file type."""
    file_type = file_type.lower()

    if file_type == "pdf" or file_type == ".pdf":
        return extract_text_from_pdf_bytes(content)
    elif file_type in ("docx", ".docx", "doc", ".doc"):
        return extract_text_from_docx_bytes(content)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


def chunk_text(text: str, chunk_size: int = 8000, overlap: int = 500) -> list[str]:
    """Split text into overlapping chunks for LLM processing."""
    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0

    while start < len(text):
        end = start + chunk_size

        # Try to break at a paragraph or sentence boundary
        if end < len(text):
            # Look for paragraph break
            para_break = text.rfind("\n\n", start, end)
            if para_break > start + chunk_size // 2:
                end = para_break + 2
            else:
                # Look for sentence break
                sentence_break = text.rfind(". ", start, end)
                if sentence_break > start + chunk_size // 2:
                    end = sentence_break + 2

        chunks.append(text[start:end].strip())
        start = end - overlap

    return chunks
