"""Script to generate synthetic legal contract test documents."""

from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors

from .config import TEST_DOCS_DIR


def create_pdf(filename: str, content_func):
    """Create a PDF document."""
    filepath = TEST_DOCS_DIR / filename
    doc = SimpleDocTemplate(
        str(filepath),
        pagesize=letter,
        rightMargin=72,
        leftMargin=72,
        topMargin=72,
        bottomMargin=72,
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name='ContractTitle',
        parent=styles['Heading1'],
        fontSize=18,
        spaceAfter=30,
        alignment=1,  # Center
    ))
    styles.add(ParagraphStyle(
        name='SectionHeader',
        parent=styles['Heading2'],
        fontSize=12,
        spaceBefore=20,
        spaceAfter=10,
        fontName='Helvetica-Bold',
    ))
    styles.add(ParagraphStyle(
        name='ContractBody',
        parent=styles['Normal'],
        fontSize=10,
        spaceBefore=6,
        spaceAfter=6,
        leading=14,
    ))

    story = content_func(styles)
    doc.build(story)
    print(f"Created: {filepath}")


def create_docx(filename: str, content_func):
    """Create a DOCX document."""
    filepath = TEST_DOCS_DIR / filename
    doc = Document()
    content_func(doc)
    doc.save(filepath)
    print(f"Created: {filepath}")


# PDF Content Functions

def service_agreement_content(styles):
    """IT Services Agreement content."""
    effective_date = date.today().replace(day=1)
    termination_date = effective_date + timedelta(days=365)

    story = []
    story.append(Paragraph("IT SERVICES AGREEMENT", styles['ContractTitle']))

    story.append(Paragraph(
        f"This IT Services Agreement (\"Agreement\") is entered into as of {effective_date.strftime('%B %d, %Y')} "
        "(\"Effective Date\") by and between:",
        styles['ContractBody']
    ))

    story.append(Spacer(1, 12))
    story.append(Paragraph(
        "<b>TechSolutions Inc.</b>, a Delaware corporation with offices at 123 Innovation Drive, San Francisco, CA 94105 "
        "(\"Service Provider\")",
        styles['ContractBody']
    ))
    story.append(Paragraph("and", styles['ContractBody']))
    story.append(Paragraph(
        "<b>GlobalRetail Corp.</b>, a California corporation with offices at 456 Commerce Street, Los Angeles, CA 90001 "
        "(\"Client\")",
        styles['ContractBody']
    ))

    story.append(Paragraph("1. SERVICES", styles['SectionHeader']))
    story.append(Paragraph(
        "Service Provider agrees to provide the following IT services to Client: (a) Cloud infrastructure management; "
        "(b) 24/7 technical support; (c) Security monitoring and incident response; (d) Monthly system maintenance and updates.",
        styles['ContractBody']
    ))

    story.append(Paragraph("2. TERM", styles['SectionHeader']))
    story.append(Paragraph(
        f"This Agreement shall commence on the Effective Date and continue for a period of twelve (12) months, "
        f"terminating on {termination_date.strftime('%B %d, %Y')}, unless earlier terminated in accordance with Section 7.",
        styles['ContractBody']
    ))

    story.append(Paragraph("3. COMPENSATION", styles['SectionHeader']))
    story.append(Paragraph(
        "Client shall pay Service Provider a monthly fee of $15,000 USD, due on the first business day of each month. "
        "Additional services shall be billed at $200 per hour.",
        styles['ContractBody']
    ))

    story.append(Paragraph("4. CONFIDENTIALITY", styles['SectionHeader']))
    story.append(Paragraph(
        "Each party agrees to maintain the confidentiality of all proprietary information received from the other party. "
        "This obligation shall survive termination of this Agreement for a period of three (3) years.",
        styles['ContractBody']
    ))

    story.append(Paragraph("5. SERVICE LEVEL AGREEMENT", styles['SectionHeader']))
    story.append(Paragraph(
        "Service Provider guarantees 99.9% uptime for all managed systems. Failure to meet this SLA shall result in "
        "service credits equal to 10% of the monthly fee for each 0.1% below the guaranteed uptime.",
        styles['ContractBody']
    ))

    story.append(Paragraph("6. LIABILITY", styles['SectionHeader']))
    story.append(Paragraph(
        "Neither party shall be liable for any indirect, incidental, or consequential damages. Service Provider's "
        "total liability shall not exceed the fees paid by Client in the twelve (12) months preceding the claim.",
        styles['ContractBody']
    ))

    story.append(Paragraph("7. TERMINATION", styles['SectionHeader']))
    story.append(Paragraph(
        "Either party may terminate this Agreement with 30 days written notice. Either party may terminate immediately "
        "upon material breach by the other party that remains uncured for 15 days after written notice.",
        styles['ContractBody']
    ))

    story.append(Spacer(1, 40))
    story.append(Paragraph("SIGNATURES", styles['SectionHeader']))
    story.append(Spacer(1, 20))

    sig_data = [
        ['TechSolutions Inc.', '', 'GlobalRetail Corp.'],
        ['_______________________', '', '_______________________'],
        ['John Smith, CEO', '', 'Maria Garcia, COO'],
        [f'Date: {effective_date.strftime("%m/%d/%Y")}', '', f'Date: {effective_date.strftime("%m/%d/%Y")}'],
    ]
    sig_table = Table(sig_data, colWidths=[2.5*inch, 1*inch, 2.5*inch])
    sig_table.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(sig_table)

    return story


def nda_mutual_content(styles):
    """Mutual Non-Disclosure Agreement content."""
    effective_date = date.today()

    story = []
    story.append(Paragraph("MUTUAL NON-DISCLOSURE AGREEMENT", styles['ContractTitle']))

    story.append(Paragraph(
        f"This Mutual Non-Disclosure Agreement (\"Agreement\") is made effective as of {effective_date.strftime('%B %d, %Y')} "
        "by and between:",
        styles['ContractBody']
    ))

    story.append(Spacer(1, 12))
    story.append(Paragraph(
        "<b>InnovateTech Labs, LLC</b>, a California limited liability company located at 789 Research Park, "
        "Palo Alto, CA 94301 (\"Party A\")",
        styles['ContractBody']
    ))
    story.append(Paragraph("and", styles['ContractBody']))
    story.append(Paragraph(
        "<b>DataDynamics Corp.</b>, a Nevada corporation located at 321 Enterprise Boulevard, "
        "Reno, NV 89501 (\"Party B\")",
        styles['ContractBody']
    ))

    story.append(Paragraph("RECITALS", styles['SectionHeader']))
    story.append(Paragraph(
        "The parties wish to explore a potential business relationship concerning joint development of AI-powered "
        "analytics solutions (\"Purpose\"). In connection with this Purpose, each party may disclose certain "
        "confidential and proprietary information to the other.",
        styles['ContractBody']
    ))

    story.append(Paragraph("1. DEFINITION OF CONFIDENTIAL INFORMATION", styles['SectionHeader']))
    story.append(Paragraph(
        "\"Confidential Information\" means any non-public information disclosed by either party, including but not "
        "limited to: trade secrets, technical data, algorithms, source code, business plans, customer lists, "
        "financial information, and marketing strategies.",
        styles['ContractBody']
    ))

    story.append(Paragraph("2. OBLIGATIONS OF RECEIVING PARTY", styles['SectionHeader']))
    story.append(Paragraph(
        "Each party agrees to: (a) hold Confidential Information in strict confidence; (b) not disclose Confidential "
        "Information to third parties without prior written consent; (c) use Confidential Information only for the Purpose; "
        "(d) protect Confidential Information with the same degree of care used for its own confidential information.",
        styles['ContractBody']
    ))

    story.append(Paragraph("3. EXCLUSIONS", styles['SectionHeader']))
    story.append(Paragraph(
        "Confidential Information does not include information that: (a) is or becomes publicly available through no "
        "fault of the receiving party; (b) was rightfully in possession of the receiving party prior to disclosure; "
        "(c) is independently developed without use of Confidential Information; (d) is rightfully obtained from a third party.",
        styles['ContractBody']
    ))

    story.append(Paragraph("4. TERM", styles['SectionHeader']))
    story.append(Paragraph(
        "This Agreement shall remain in effect for two (2) years from the Effective Date. The confidentiality "
        "obligations shall survive termination for an additional three (3) years.",
        styles['ContractBody']
    ))

    story.append(Paragraph("5. RETURN OF MATERIALS", styles['SectionHeader']))
    story.append(Paragraph(
        "Upon termination or request, each party shall promptly return or destroy all Confidential Information "
        "received from the other party, and certify such destruction in writing.",
        styles['ContractBody']
    ))

    story.append(Paragraph("6. REMEDIES", styles['SectionHeader']))
    story.append(Paragraph(
        "Each party acknowledges that breach of this Agreement may cause irreparable harm for which monetary damages "
        "would be inadequate. The non-breaching party shall be entitled to seek injunctive relief in addition to any "
        "other remedies available at law or in equity.",
        styles['ContractBody']
    ))

    story.append(Spacer(1, 30))
    story.append(Paragraph("IN WITNESS WHEREOF, the parties have executed this Agreement:", styles['ContractBody']))
    story.append(Spacer(1, 20))

    sig_data = [
        ['InnovateTech Labs, LLC', '', 'DataDynamics Corp.'],
        ['_______________________', '', '_______________________'],
        ['Dr. Sarah Chen, Managing Partner', '', 'Robert Williams, President'],
        [f'Date: {effective_date.strftime("%m/%d/%Y")}', '', f'Date: {effective_date.strftime("%m/%d/%Y")}'],
    ]
    sig_table = Table(sig_data, colWidths=[2.5*inch, 1*inch, 2.5*inch])
    sig_table.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
    ]))
    story.append(sig_table)

    return story


def lease_commercial_content(styles):
    """Commercial Property Lease Agreement content."""
    effective_date = date.today().replace(day=1)
    termination_date = effective_date + timedelta(days=365*3)

    story = []
    story.append(Paragraph("COMMERCIAL LEASE AGREEMENT", styles['ContractTitle']))

    story.append(Paragraph(
        f"This Commercial Lease Agreement (\"Lease\") is entered into as of {effective_date.strftime('%B %d, %Y')} "
        "between:",
        styles['ContractBody']
    ))

    story.append(Spacer(1, 12))
    story.append(Paragraph(
        "<b>Metropolitan Properties LLC</b>, a Texas limited liability company with principal offices at "
        "1000 Skyline Tower, Houston, TX 77002 (\"Landlord\")",
        styles['ContractBody']
    ))
    story.append(Paragraph("and", styles['ContractBody']))
    story.append(Paragraph(
        "<b>StartupHub Ventures Inc.</b>, a Delaware corporation with offices at "
        "555 Innovation Way, Austin, TX 78701 (\"Tenant\")",
        styles['ContractBody']
    ))

    story.append(Paragraph("1. PREMISES", styles['SectionHeader']))
    story.append(Paragraph(
        "Landlord leases to Tenant the commercial space located at Suite 400, 2500 Commerce Center Drive, "
        "Houston, TX 77003, consisting of approximately 5,000 square feet of office space (\"Premises\").",
        styles['ContractBody']
    ))

    story.append(Paragraph("2. TERM", styles['SectionHeader']))
    story.append(Paragraph(
        f"The lease term shall commence on {effective_date.strftime('%B %d, %Y')} and terminate on "
        f"{termination_date.strftime('%B %d, %Y')}, a period of three (3) years, unless renewed or terminated earlier.",
        styles['ContractBody']
    ))

    story.append(Paragraph("3. RENT", styles['SectionHeader']))
    story.append(Paragraph(
        "Tenant shall pay monthly rent of $12,500 USD, due on the first day of each month. Rent shall increase "
        "by 3% annually on each anniversary of the Effective Date. Late payments shall incur a fee of 5% of the monthly rent.",
        styles['ContractBody']
    ))

    story.append(Paragraph("4. SECURITY DEPOSIT", styles['SectionHeader']))
    story.append(Paragraph(
        "Tenant shall deposit $25,000 USD as security deposit upon execution of this Lease. The deposit shall be "
        "returned within 30 days of lease termination, less any amounts owed for damages or unpaid rent.",
        styles['ContractBody']
    ))

    story.append(Paragraph("5. USE OF PREMISES", styles['SectionHeader']))
    story.append(Paragraph(
        "Tenant shall use the Premises solely for general office purposes, technology development, and related "
        "lawful business activities. No manufacturing, retail sales, or hazardous materials storage is permitted.",
        styles['ContractBody']
    ))

    story.append(Paragraph("6. MAINTENANCE AND REPAIRS", styles['SectionHeader']))
    story.append(Paragraph(
        "Landlord shall maintain the building structure, HVAC systems, and common areas. Tenant shall maintain "
        "the interior of the Premises in good condition and repair any damage caused by Tenant or its agents.",
        styles['ContractBody']
    ))

    story.append(Paragraph("7. INSURANCE", styles['SectionHeader']))
    story.append(Paragraph(
        "Tenant shall maintain commercial general liability insurance with minimum coverage of $1,000,000 per occurrence "
        "and $2,000,000 aggregate, naming Landlord as additional insured.",
        styles['ContractBody']
    ))

    story.append(Paragraph("8. TERMINATION", styles['SectionHeader']))
    story.append(Paragraph(
        "Either party may terminate upon 90 days written notice for material breach that remains uncured for 30 days. "
        "Tenant may not terminate early without paying a penalty equal to three months' rent.",
        styles['ContractBody']
    ))

    story.append(Spacer(1, 30))
    sig_data = [
        ['LANDLORD:', '', 'TENANT:'],
        ['Metropolitan Properties LLC', '', 'StartupHub Ventures Inc.'],
        ['_______________________', '', '_______________________'],
        ['James Mitchell, Managing Member', '', 'Emily Rodriguez, CEO'],
        [f'Date: {effective_date.strftime("%m/%d/%Y")}', '', f'Date: {effective_date.strftime("%m/%d/%Y")}'],
    ]
    sig_table = Table(sig_data, colWidths=[2.5*inch, 1*inch, 2.5*inch])
    sig_table.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
    ]))
    story.append(sig_table)

    return story


# DOCX Content Functions

def employment_contract_content(doc):
    """Employment Agreement content."""
    effective_date = date.today()

    title = doc.add_heading('EMPLOYMENT AGREEMENT', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        f"This Employment Agreement (\"Agreement\") is made effective as of {effective_date.strftime('%B %d, %Y')} "
        f"by and between:"
    )

    doc.add_paragraph(
        "Quantum Software Solutions, Inc., a Massachusetts corporation with headquarters at "
        "100 Technology Square, Cambridge, MA 02139 (\"Employer\")"
    )
    doc.add_paragraph(
        "and Michael Thomas Anderson, residing at 45 Beacon Street, Boston, MA 02108 (\"Employee\")"
    )

    doc.add_heading('1. POSITION AND DUTIES', level=1)
    doc.add_paragraph(
        "Employee is hired as Senior Software Engineer, reporting to the Chief Technology Officer. "
        "Employee's duties shall include: software development, code review, mentoring junior developers, "
        "and participating in architectural decisions. Employee agrees to devote full professional time "
        "and attention to these duties."
    )

    doc.add_heading('2. TERM OF EMPLOYMENT', level=1)
    doc.add_paragraph(
        "This Agreement shall commence on the Effective Date and continue for an indefinite period until "
        "terminated by either party in accordance with Section 8."
    )

    doc.add_heading('3. COMPENSATION', level=1)
    doc.add_paragraph(
        "Employee shall receive an annual base salary of $185,000 USD, payable in bi-weekly installments. "
        "Salary shall be reviewed annually and may be adjusted at Employer's discretion."
    )

    doc.add_heading('4. BENEFITS', level=1)
    benefits = doc.add_paragraph()
    benefits.add_run("Employee shall be entitled to the following benefits:\n")
    benefits.add_run("• Health, dental, and vision insurance (premium paid 80% by Employer)\n")
    benefits.add_run("• 401(k) retirement plan with 4% employer match\n")
    benefits.add_run("• 20 days paid vacation annually\n")
    benefits.add_run("• 10 days paid sick leave annually\n")
    benefits.add_run("• Stock options: 10,000 shares vesting over 4 years\n")
    benefits.add_run("• Annual professional development budget of $5,000")

    doc.add_heading('5. CONFIDENTIALITY', level=1)
    doc.add_paragraph(
        "Employee agrees to maintain confidentiality of all proprietary information, trade secrets, "
        "and business processes. This obligation survives termination of employment indefinitely for "
        "trade secrets and for three (3) years for other confidential information."
    )

    doc.add_heading('6. NON-COMPETE', level=1)
    doc.add_paragraph(
        "For twelve (12) months following termination, Employee shall not work for any direct competitor "
        "in the enterprise software sector within a 50-mile radius of Employer's offices, nor solicit "
        "any of Employer's clients or employees."
    )

    doc.add_heading('7. INTELLECTUAL PROPERTY', level=1)
    doc.add_paragraph(
        "All inventions, software, and intellectual property created by Employee during employment "
        "and related to Employer's business shall be the sole property of Employer."
    )

    doc.add_heading('8. TERMINATION', level=1)
    doc.add_paragraph(
        "Either party may terminate this Agreement with two (2) weeks written notice. Employer may "
        "terminate immediately for cause, including material breach, misconduct, or failure to perform duties. "
        "Upon termination, Employee shall return all company property and confidential materials."
    )

    doc.add_paragraph("")
    doc.add_paragraph("AGREED AND ACCEPTED:")
    doc.add_paragraph("")

    # Signature table
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "EMPLOYER:"
    table.cell(0, 2).text = "EMPLOYEE:"
    table.cell(1, 0).text = "Quantum Software Solutions, Inc."
    table.cell(1, 2).text = "Michael Thomas Anderson"
    table.cell(2, 0).text = f"By: Dr. Patricia Lee, CEO\nDate: {effective_date.strftime('%m/%d/%Y')}"
    table.cell(2, 2).text = f"Date: {effective_date.strftime('%m/%d/%Y')}"


def partnership_agreement_content(doc):
    """Business Partnership Agreement content."""
    effective_date = date.today().replace(day=1)

    title = doc.add_heading('PARTNERSHIP AGREEMENT', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        f"This Partnership Agreement (\"Agreement\") is entered into as of {effective_date.strftime('%B %d, %Y')} "
        f"by and among the following Partners:"
    )

    doc.add_paragraph("1. Alexandra Chen, an individual residing at 234 Oak Avenue, Seattle, WA 98101")
    doc.add_paragraph("2. Benjamin Harris, an individual residing at 567 Pine Street, Seattle, WA 98102")
    doc.add_paragraph("3. Catherine Davis, an individual residing at 890 Maple Lane, Bellevue, WA 98004")

    doc.add_paragraph(
        "The Partners hereby agree to form a general partnership under the name \"Pacific Consulting Group\" "
        "(\"Partnership\") for the purpose of providing business consulting services."
    )

    doc.add_heading('1. FORMATION AND NAME', level=1)
    doc.add_paragraph(
        "The Partnership is formed under the laws of the State of Washington. The principal place of business "
        "shall be 1200 First Avenue, Suite 500, Seattle, WA 98101."
    )

    doc.add_heading('2. CAPITAL CONTRIBUTIONS', level=1)
    contributions = doc.add_paragraph()
    contributions.add_run("Each Partner shall contribute the following initial capital:\n")
    contributions.add_run("• Alexandra Chen: $100,000 (40% ownership)\n")
    contributions.add_run("• Benjamin Harris: $75,000 (30% ownership)\n")
    contributions.add_run("• Catherine Davis: $75,000 (30% ownership)\n")
    contributions.add_run("\nAdditional capital contributions require unanimous consent of all Partners.")

    doc.add_heading('3. PROFIT AND LOSS ALLOCATION', level=1)
    doc.add_paragraph(
        "Profits and losses shall be allocated in proportion to each Partner's ownership percentage. "
        "Distributions shall be made quarterly, within 30 days of quarter end, after retaining sufficient "
        "working capital as determined by majority vote."
    )

    doc.add_heading('4. MANAGEMENT AND DUTIES', level=1)
    doc.add_paragraph(
        "All Partners shall have equal rights in management of the Partnership. Major decisions, including "
        "expenditures over $25,000, hiring of employees, and entering contracts over $50,000, require "
        "majority approval. Each Partner agrees to devote necessary time and attention to Partnership business."
    )

    doc.add_heading('5. PARTNER COMPENSATION', level=1)
    doc.add_paragraph(
        "Each Partner shall receive a monthly draw of $8,000 against their share of profits. Additional "
        "compensation shall be by unanimous agreement only."
    )

    doc.add_heading('6. WITHDRAWAL OF A PARTNER', level=1)
    doc.add_paragraph(
        "A Partner may withdraw with 90 days written notice. Upon withdrawal, the withdrawing Partner shall "
        "be entitled to their capital account balance, payable over 24 months. The remaining Partners shall "
        "have first right to purchase the withdrawing Partner's interest at fair market value."
    )

    doc.add_heading('7. NON-COMPETE AND CONFIDENTIALITY', level=1)
    doc.add_paragraph(
        "During participation and for two (2) years after withdrawal, no Partner shall engage in competing "
        "business consulting services in the Pacific Northwest region. All Partners agree to maintain "
        "confidentiality of Partnership business information."
    )

    doc.add_heading('8. DISSOLUTION', level=1)
    doc.add_paragraph(
        "The Partnership shall dissolve upon: (a) unanimous agreement of Partners; (b) death or incapacity of "
        "a Partner, unless remaining Partners elect to continue; (c) bankruptcy of the Partnership. Upon "
        "dissolution, assets shall be liquidated and distributed first to creditors, then to Partners."
    )

    doc.add_paragraph("")
    doc.add_paragraph(
        f"IN WITNESS WHEREOF, the Partners have executed this Agreement as of {effective_date.strftime('%B %d, %Y')}."
    )
    doc.add_paragraph("")

    for name in ["Alexandra Chen", "Benjamin Harris", "Catherine Davis"]:
        doc.add_paragraph(f"________________________\n{name}\nDate: {effective_date.strftime('%m/%d/%Y')}")
        doc.add_paragraph("")


def consulting_agreement_content(doc):
    """Independent Contractor/Consulting Agreement content."""
    effective_date = date.today()
    termination_date = effective_date + timedelta(days=180)

    title = doc.add_heading('CONSULTING AGREEMENT', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        f"This Consulting Agreement (\"Agreement\") is made effective as of {effective_date.strftime('%B %d, %Y')} "
        f"by and between:"
    )

    doc.add_paragraph(
        "FinanceFirst Technologies, Inc., a New York corporation with offices at 350 Fifth Avenue, "
        "New York, NY 10118 (\"Company\")"
    )
    doc.add_paragraph("and")
    doc.add_paragraph(
        "David Morrison, an independent consultant with a principal address at 789 Riverside Drive, "
        "New York, NY 10032 (\"Consultant\")"
    )

    doc.add_heading('RECITALS', level=1)
    doc.add_paragraph(
        "Company desires to engage Consultant to provide expert financial systems consulting services, "
        "and Consultant desires to provide such services on the terms set forth herein."
    )

    doc.add_heading('1. ENGAGEMENT', level=1)
    doc.add_paragraph(
        "Company engages Consultant to provide financial systems integration consulting, including: "
        "system analysis, architecture design, implementation oversight, and staff training for "
        "Company's new enterprise resource planning (ERP) system."
    )

    doc.add_heading('2. TERM', level=1)
    doc.add_paragraph(
        f"This Agreement shall commence on {effective_date.strftime('%B %d, %Y')} and terminate on "
        f"{termination_date.strftime('%B %d, %Y')}, unless extended by mutual written agreement or "
        f"terminated earlier in accordance with Section 7."
    )

    doc.add_heading('3. COMPENSATION', level=1)
    comp = doc.add_paragraph()
    comp.add_run("Consultant shall be compensated as follows:\n")
    comp.add_run("• Hourly rate: $275 USD per hour\n")
    comp.add_run("• Maximum weekly hours: 40 hours\n")
    comp.add_run("• Minimum monthly retainer: $20,000 USD\n")
    comp.add_run("• Expenses: Pre-approved travel and materials reimbursed at cost\n")
    comp.add_run("\nInvoices shall be submitted bi-weekly with payment due within 15 days.")

    doc.add_heading('4. INDEPENDENT CONTRACTOR STATUS', level=1)
    doc.add_paragraph(
        "Consultant is an independent contractor, not an employee. Consultant is responsible for all taxes, "
        "insurance, and benefits. Consultant may provide services to other clients provided there is no "
        "conflict of interest with Company's business."
    )

    doc.add_heading('5. DELIVERABLES', level=1)
    deliverables = doc.add_paragraph()
    deliverables.add_run("Consultant shall deliver:\n")
    deliverables.add_run("• System requirements document (within 30 days)\n")
    deliverables.add_run("• Architecture design specification (within 60 days)\n")
    deliverables.add_run("• Implementation plan (within 90 days)\n")
    deliverables.add_run("• Training materials (within 120 days)\n")
    deliverables.add_run("• Final project report (within 180 days)")

    doc.add_heading('6. CONFIDENTIALITY', level=1)
    doc.add_paragraph(
        "Consultant agrees to maintain strict confidentiality of all Company information, including financial "
        "data, business strategies, and proprietary systems. This obligation survives termination for five (5) years."
    )

    doc.add_heading('7. TERMINATION', level=1)
    doc.add_paragraph(
        "Either party may terminate with 14 days written notice. Company may terminate immediately for cause. "
        "Upon termination, Consultant shall deliver all work product and be paid for services rendered through "
        "the termination date."
    )

    doc.add_heading('8. INTELLECTUAL PROPERTY', level=1)
    doc.add_paragraph(
        "All work product, including documents, designs, and code, created under this Agreement shall be the "
        "exclusive property of Company. Consultant retains rights to general knowledge, skills, and tools "
        "that existed prior to the engagement."
    )

    doc.add_heading('9. LIABILITY', level=1)
    doc.add_paragraph(
        "Consultant's liability shall be limited to direct damages and shall not exceed the total fees paid "
        "under this Agreement. Neither party shall be liable for indirect or consequential damages."
    )

    doc.add_paragraph("")
    doc.add_paragraph("AGREED AND ACCEPTED:")
    doc.add_paragraph("")

    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "COMPANY:"
    table.cell(0, 2).text = "CONSULTANT:"
    table.cell(1, 0).text = "FinanceFirst Technologies, Inc."
    table.cell(1, 2).text = "David Morrison"
    table.cell(2, 0).text = f"By: Jennifer Walsh, VP Operations\nDate: {effective_date.strftime('%m/%d/%Y')}"
    table.cell(2, 2).text = f"Date: {effective_date.strftime('%m/%d/%Y')}"


def generate_all_test_documents():
    """Generate all synthetic test documents."""
    TEST_DOCS_DIR.mkdir(exist_ok=True)

    print("Generating synthetic legal contract test documents...")

    # PDFs
    create_pdf("service_agreement.pdf", service_agreement_content)
    create_pdf("nda_mutual.pdf", nda_mutual_content)
    create_pdf("lease_commercial.pdf", lease_commercial_content)

    # DOCX
    create_docx("employment_contract.docx", employment_contract_content)
    create_docx("partnership_agreement.docx", partnership_agreement_content)
    create_docx("consulting_agreement.docx", consulting_agreement_content)

    print(f"\nAll test documents created in: {TEST_DOCS_DIR}")


if __name__ == "__main__":
    generate_all_test_documents()
